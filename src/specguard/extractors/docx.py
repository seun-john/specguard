"""DOCX extractor built on python-docx.

Covers body paragraphs, headings, list items, tables, section geometry and a few
document defaults. It does not lay out pages, and it ignores text boxes, headers,
footers, footnotes and comments.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from specguard.errors import DocumentReadError
from specguard.extractors.base import DocumentExtractor, register_extractor
from specguard.models.document import Document, Heading, Paragraph, Table

# Refuse archives that would expand far beyond a plausible document (zip bombs).
MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 5000

_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)


def _check_archive(data: bytes, source: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            members = zf.infolist()
    except zipfile.BadZipFile:
        raise DocumentReadError(f"{source} is not a valid .docx file (not a ZIP archive)") from None
    if len(members) > MAX_ARCHIVE_MEMBERS:
        raise DocumentReadError(f"{source} contains too many archive entries to be a normal .docx")
    if sum(m.file_size for m in members) > MAX_UNCOMPRESSED_BYTES:
        raise DocumentReadError(f"{source} expands to an unreasonable size and was not opened")
    if "word/document.xml" not in {m.filename for m in members}:
        raise DocumentReadError(f"{source} is a ZIP archive but not a Word document")


def _heading_level(par: DocxParagraph) -> int | None:
    """Heading level from the style, or from an outline level set on the paragraph."""
    style = par.style
    names = [getattr(style, "name", None) or "", getattr(style, "style_id", None) or ""]
    for name in names:
        m = _HEADING_STYLE_RE.match(name)
        if m:
            return int(m.group(1))
        if name.lower() == "title":
            return 1
    outline = par._p.xpath("./w:pPr/w:outlineLvl/@w:val")
    if outline:
        try:
            return int(outline[0]) + 1
        except ValueError:
            return None
    return None


def _list_marker(par: DocxParagraph, numbering: Any) -> str | None:
    """`bullet`, `numbered` or `unknown` for list paragraphs, None for others."""
    style_name = (getattr(par.style, "name", None) or "").lower()
    if style_name.startswith("list bullet"):
        return "bullet"
    if style_name.startswith("list number"):
        return "numbered"
    num_ids = par._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
    if not num_ids:
        return None
    try:
        ilvl = (par._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val") or ["0"])[0]
        abstract = numbering.xpath(f'./w:num[@w:numId="{num_ids[0]}"]/w:abstractNumId/@w:val')
        fmt = numbering.xpath(
            f'./w:abstractNum[@w:abstractNumId="{abstract[0]}"]'
            f'/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val'
        )
        return "bullet" if fmt and fmt[0] == "bullet" else "numbered"
    except (IndexError, AttributeError):
        return "unknown"


def _length(value: Any, unit: str = "cm") -> float | None:
    return None if value is None else round(float(getattr(value, unit)), 2)


def _collect_metadata(document: Any) -> dict[str, Any]:
    meta: dict[str, Any] = {}
    core = document.core_properties
    meta["title"] = core.title or None
    meta["author"] = core.author or None
    meta["last_modified_by"] = core.last_modified_by or None
    meta["created"] = core.created.isoformat() if core.created else None
    meta["modified"] = core.modified.isoformat() if core.modified else None

    sections = []
    for s in document.sections:
        sections.append(
            {
                "orientation": "landscape" if s.orientation == 1 else "portrait",
                "page_width_cm": _length(s.page_width),
                "page_height_cm": _length(s.page_height),
                "margin_top_cm": _length(s.top_margin),
                "margin_bottom_cm": _length(s.bottom_margin),
                "margin_left_cm": _length(s.left_margin),
                "margin_right_cm": _length(s.right_margin),
            }
        )
    meta["sections"] = sections

    try:
        normal = document.styles["Normal"]
        meta["default_font_name"] = normal.font.name
        meta["default_font_size_pt"] = _length(normal.font.size, "pt")
        spacing = normal.paragraph_format.line_spacing
        if spacing is None:
            meta["default_line_spacing"] = None
        elif isinstance(spacing, float):
            meta["default_line_spacing"] = round(spacing, 2)
        else:
            meta["default_line_spacing_pt"] = _length(spacing, "pt")
            meta["default_line_spacing"] = None
    except KeyError:
        meta["default_font_name"] = None
    return meta


class DocxExtractor(DocumentExtractor):
    """`.docx` files. Paragraph numbers count non-empty body paragraphs in order."""

    file_type = "docx"
    extensions = (".docx",)

    def extract(self, data: bytes, source: str) -> Document:
        _check_archive(data, source)
        try:
            word = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise DocumentReadError(
                f"{source} could not be opened as a Word document ({type(exc).__name__})"
            ) from exc

        try:
            numbering = word.part.numbering_part.element
        except Exception:  # no numbering part: no lists
            numbering = None

        doc = Document(file_type="docx", raw_text="")
        position = 0
        para_index = 0
        texts: list[str] = []

        for child in word.element.body.iterchildren():
            if child.tag == qn("w:p"):
                par = DocxParagraph(child, word)
                text = par.text.strip()
                if not text:
                    continue
                position += 1
                para_index += 1
                level = _heading_level(par)
                marker = (
                    _list_marker(par, numbering)
                    if (level is None and numbering is not None)
                    else None
                )
                p = Paragraph(
                    index=para_index,
                    text=text,
                    position=position,
                    kind="heading" if level else ("list_item" if marker else "paragraph"),
                    heading_level=level,
                    list_marker=marker,
                    style=getattr(par.style, "name", None),
                )
                doc.paragraphs.append(p)
                texts.append(text)
                if level:
                    doc.headings.append(
                        Heading(text=text, level=level, paragraph=p.index, position=position)
                    )
            elif child.tag == qn("w:tbl"):
                table = DocxTable(child, word)
                rows: list[list[str]] = []
                for row in table.rows:
                    seen: set[int] = set()
                    cells: list[str] = []
                    for cell in row.cells:
                        if id(cell._tc) in seen:  # merged cells repeat the same element
                            continue
                        seen.add(id(cell._tc))
                        cells.append(cell.text.strip())
                    rows.append(cells)
                position += 1
                doc.tables.append(Table(index=len(doc.tables) + 1, rows=rows, position=position))
                texts.extend(c for r in rows for c in r if c)

        doc.raw_text = "\n\n".join(texts)
        doc.metadata.update(_collect_metadata(word))
        return doc


register_extractor(DocxExtractor())

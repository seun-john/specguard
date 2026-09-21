"""The shipped examples, and the end-to-end workflow from the project brief."""

from __future__ import annotations

import json
from pathlib import Path

import docx
import pytest
from tests.helpers import FIXTURES
from typer.testing import CliRunner

from specguard.api import audit_file
from specguard.cli import app
from specguard.models.result import Status

ROOT = Path(__file__).parents[2]
EXAMPLES = ROOT / "examples"
runner = CliRunner()


def statuses(doc: Path, spec: Path) -> dict[str, Status]:
    return {e.requirement.id: e.result.status for e in audit_file(doc, spec).entries}


class TestShippedExamples:
    def test_every_example_specification_is_valid(self) -> None:
        for spec in EXAMPLES.glob("*/specguard.yml"):
            result = runner.invoke(app, ["validate", str(spec)])
            assert result.exit_code == 0, (spec, result.stderr)

    def test_academic_markdown(self) -> None:
        got = statuses(EXAMPLES / "academic" / "draft.md", EXAMPLES / "academic" / "specguard.yml")
        assert got == {
            "AC001": Status.PASS,
            "AC002": Status.FAIL,  # em dash in Methodology
            "AC003": Status.FAIL,  # no Conclusion
            "AC004": Status.FAIL,  # first research question was reworded
            "AC005": Status.FAIL,  # "It is important to note that"
            "AC006": Status.FAIL,  # 3 references, 5 required
            "AC007": Status.UNVERIFIED,
            "AC008": Status.UNVERIFIED,
        }

    def test_academic_docx(self) -> None:
        got = statuses(
            EXAMPLES / "academic" / "thesis_sample.docx", EXAMPLES / "academic" / "specguard.yml"
        )
        assert got["AC003"] is Status.PASS
        assert got["AC004"] is Status.PASS  # questions preserved in the Word version
        assert got["AC002"] is Status.FAIL
        assert got["AC006"] is Status.FAIL

    def test_the_docx_example_is_a_real_word_file_and_is_small(self) -> None:
        path = EXAMPLES / "academic" / "thesis_sample.docx"
        assert path.stat().st_size < 60_000
        assert docx.Document(str(path)).paragraphs

    def test_business(self) -> None:
        got = statuses(
            EXAMPLES / "business" / "quarterly_report.md", EXAMPLES / "business" / "specguard.yml"
        )
        assert got == {
            "BR001": Status.PASS,
            "BR002": Status.PASS,
            "BR003": Status.FAIL,  # "Next Steps" is not a Recommendations heading
            "BR004": Status.FAIL,  # TBD
            "BR005": Status.FAIL,  # [Insert mitigation owner]
            "BR006": Status.UNVERIFIED,
        }

    def test_coding(self) -> None:
        got = statuses(
            EXAMPLES / "coding" / "README.example.md", EXAMPLES / "coding" / "specguard.yml"
        )
        assert got == {
            "RD001": Status.PASS,
            "RD002": Status.PASS,
            "RD003": Status.PASS,
            "RD004": Status.FAIL,  # TODO
            "RD005": Status.PASS,
        }

    def test_the_examples_cover_more_than_one_domain(self) -> None:
        assert {p.parent.name for p in EXAMPLES.glob("*/specguard.yml")} == {
            "academic",
            "business",
            "coding",
        }


class TestBriefToAuditWorkflow:
    """brief.txt -> specguard extract -> specguard audit report.docx."""

    def make_report(self, path: Path, *, words: int, em_dashes: int, headings: list[str]) -> Path:
        d = docx.Document()
        d.add_heading("Approved Title", 0)
        filler = ("data " * 40).strip() + "."
        for name in headings:
            d.add_heading(name, 1)
            d.add_paragraph(filler)
        remaining = words - 1 - len(headings) - 40 * len(headings)
        d.add_paragraph(" ".join(["more"] * max(remaining, 0)) + ".")
        for _ in range(em_dashes):
            d.add_paragraph("Something — else.")
        d.save(str(path))
        return path

    def test_full_workflow(self, tmp_path: Path) -> None:
        brief = FIXTURES / "instructions.txt"
        spec = tmp_path / "specguard.yml"
        assert runner.invoke(app, ["extract", str(brief), "-o", str(spec)]).exit_code == 0

        short = self.make_report(
            tmp_path / "report.docx",
            words=1200,
            em_dashes=2,
            headings=["Background", "Findings", "Conclusion"],
        )
        result = runner.invoke(app, ["audit", str(short), "--spec", str(spec), "--format", "json"])
        assert result.exit_code == 1
        report = json.loads(result.stdout)
        by_id = {e["requirement"]["id"]: e for e in report["entries"]}

        # SG001 minimum words FAIL, SG002 maximum PASS, SG003 em dashes FAIL
        assert by_id["SG001"]["result"]["status"] == "FAIL"
        assert "1,5" in by_id["SG001"]["result"]["expected"]
        assert by_id["SG001"]["result"]["details"]["word_count"] < 1500
        assert by_id["SG002"]["result"]["status"] == "PASS"
        assert by_id["SG003"]["result"]["status"] == "FAIL"
        assert by_id["SG003"]["result"]["details"]["occurrences"] == 2
        assert len(by_id["SG003"]["result"]["locations"]) == 2
        # Recommendations is missing; Background, Findings, Conclusion are present
        assert [by_id[i]["result"]["status"] for i in ("SG004", "SG005", "SG006", "SG007")] == [
            "PASS",
            "PASS",
            "FAIL",
            "PASS",
        ]
        # title (needs the protected text) and style (subjective) are UNVERIFIED, never PASS
        assert by_id["SG008"]["result"]["status"] == "UNVERIFIED"
        assert by_id["SG009"]["result"]["status"] == "UNVERIFIED"
        summary = report["summary"]
        assert summary["counts"] == {
            "PASS": 4,
            "FAIL": 3,
            "WARNING": 0,
            "UNVERIFIED": 2,
            "ERROR": 0,
        }
        assert summary["verification_coverage"] == pytest.approx(77.8)
        assert summary["compliance_score"] is not None and summary["compliance_score"] < 100

    def test_a_conforming_report_passes_every_testable_rule(self, tmp_path: Path) -> None:
        spec = tmp_path / "specguard.yml"
        runner.invoke(app, ["extract", str(FIXTURES / "instructions.txt"), "-o", str(spec)])
        good = self.make_report(
            tmp_path / "good.docx",
            words=1700,
            em_dashes=0,
            headings=["Background", "Findings", "Recommendations", "Conclusion"],
        )
        report = audit_file(good, spec)
        assert report.summary.counts["FAIL"] == 0
        assert report.summary.counts["PASS"] == 7
        assert report.summary.counts["UNVERIFIED"] == 2  # still not silently passed

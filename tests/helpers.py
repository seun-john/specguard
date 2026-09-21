"""Shared test helpers."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from specguard.extractors import load_document_from_text
from specguard.models.document import Document
from specguard.models.requirement import Requirement, Severity, VerificationType
from specguard.models.result import CheckResult
from specguard.rules import AuditContext, default_registry

FIXTURES = Path(__file__).parent / "fixtures"


def md(text: str) -> Document:
    """Parse Markdown text."""
    return load_document_from_text(text, "md")


def txt(text: str) -> Document:
    """Parse plain text."""
    return load_document_from_text(text, "txt")


def run_rule(
    document: Document,
    checker: str,
    parameters: dict[str, Any] | None = None,
    *,
    severity: Severity = Severity.MAJOR,
    base_dir: Path | None = None,
) -> CheckResult:
    """Run one checker directly (no engine) and return its result."""
    requirement = Requirement(
        id="T1",
        description="test requirement",
        severity=severity,
        verification_type=VerificationType.DETERMINISTIC,
        checker=checker,
        parameters=parameters or {},
    )
    instance = default_registry.create(checker, AuditContext(base_dir=base_dir))
    problems = instance.validate_parameters(requirement.parameters)
    assert not problems, problems
    return instance.check(document, requirement)


def build_docx(
    path: Path,
    *,
    landscape: bool = False,
    margin_cm: float | None = None,
    font: str | None = None,
    font_size: float | None = None,
) -> Path:
    """Write a small thesis-style DOCX with headings, a list, a table and references."""
    d = docx.Document()
    if font:
        d.styles["Normal"].font.name = font
    if font_size:
        d.styles["Normal"].font.size = Pt(font_size)
    for section in d.sections:
        if margin_cm is not None:
            section.top_margin = section.bottom_margin = Cm(margin_cm)
            section.left_margin = section.right_margin = Cm(margin_cm)
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width

    d.add_heading("Clinic Attendance Study", 0)
    d.add_heading("Introduction", 1)
    d.add_paragraph("This study examines clinic attendance in three districts.")
    d.add_paragraph("It is a short sample used for testing — nothing more.")
    d.add_heading("Methodology", 1)
    d.add_heading("Data sources", 2)
    d.add_paragraph("What factors are associated with clinic attendance?")
    d.add_paragraph("Interviews were recorded and transcribed.", style="List Bullet")
    d.add_paragraph("Records were double checked.", style="List Bullet")
    table = d.add_table(rows=2, cols=3)
    for r, row in enumerate([["District", "Visits", "Notes"], ["North", "120", "TBD figure"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_heading("Conclusion", 1)
    d.add_paragraph("Attendance improved where transport support existed.")
    d.add_heading("References", 1)
    d.add_paragraph("Adeyemi, T. (2024). Clinic access. Journal of Public Health, 12(3), 45-60.")
    d.add_paragraph("Okafor, N. (2023). Transport costs. Health Policy Review, 8(1), 12-30.")
    d.save(str(path))
    return path
